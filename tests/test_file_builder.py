"""Office file creation (plan v2 §5.1): build each format, then read it back."""

import io

import pytest

from src.file_builder import FileSpecError, build_file, safe_filename

DOC_SPEC = {
    "title": "Q3 report",
    "blocks": [
        {"type": "heading", "text": "Summary", "level": 1},
        {"type": "paragraph", "text": "Revenue grew 12%."},
        {"type": "bullets", "items": ["North +20%", "South +4%"]},
        {"type": "table", "rows": [["Region", "Revenue"], ["North", 120], ["South", 95]]},
    ],
}


def test_xlsx_round_trip_keeps_numbers_and_sheets():
    from openpyxl import load_workbook

    built = build_file("xlsx", {"sheets": [
        {"name": "Sales", "rows": [["Region", "Revenue"], ["North", 120], ["South", 95.5]]},
        {"name": "Sales", "rows": [["dup name"]]},
        {"name": "bad/name[1]", "rows": []},
    ]})
    wb = load_workbook(io.BytesIO(built.data))
    assert wb.sheetnames == ["Sales", "Sales_2", "badname1"]
    ws = wb["Sales"]
    assert [c.value for c in ws[2]] == ["North", 120]
    assert ws["B3"].value == 95.5
    assert ws["A1"].font.bold and ws.freeze_panes == "A2"


def test_spreadsheet_text_never_becomes_a_formula():
    from openpyxl import load_workbook

    rows = [["note"], ["=HYPERLINK(\"http://evil\",\"x\")"], ["+cmd"], ["-2+3"], ["@SUM(A1)"], [-5]]
    ws = load_workbook(io.BytesIO(build_file("xlsx", {"rows": rows}).data)).active
    assert ws["A2"].data_type != "f" and ws["A2"].value.startswith("'=")
    assert ws["A3"].value == "'+cmd" and ws["A4"].value == "'-2+3" and ws["A5"].value == "'@SUM(A1)"
    assert ws["A6"].value == -5  # real numbers stay numbers
    csv_text = build_file("csv", {"rows": rows}).data.decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and "\n-5" in csv_text.replace("\r", "")


def test_csv_is_utf8_with_bom_for_excel():
    data = build_file("csv", {"rows": [["name"], ["Zoë"]]}).data
    assert data.startswith(b"\xef\xbb\xbf") and "Zoë".encode() in data


def test_docx_round_trip():
    from docx import Document

    doc = Document(io.BytesIO(build_file("docx", DOC_SPEC).data))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Q3 report" in text and "Revenue grew 12%." in text and "North +20%" in text
    assert doc.tables[0].cell(1, 1).text == "120"


def test_pdf_round_trip():
    from pypdf import PdfReader

    text = PdfReader(io.BytesIO(build_file("pdf", DOC_SPEC).data)).pages[0].extract_text()
    for expected in ("Q3 report", "Summary", "Revenue grew 12%", "North", "120"):
        assert expected in text


def test_pdf_escapes_markup_in_text():
    from pypdf import PdfReader

    spec = {"blocks": [{"type": "paragraph", "text": "a < b & <b>not bold</b>"}]}
    text = PdfReader(io.BytesIO(build_file("pdf", spec).data)).pages[0].extract_text()
    assert "<b>not bold</b>" in text


def test_pptx_from_slides_and_from_blocks():
    from pptx import Presentation

    prs = Presentation(io.BytesIO(build_file("pptx", {
        "title": "Kickoff", "slides": [
            {"title": "Goals", "bullets": ["Ship v1", "Hire 2"], "notes": "say hi"},
            {"title": "Budget", "table": [["Item", "Cost"], ["GPU", 5000]]},
        ]}).data))
    titles = [s.shapes.title.text for s in prs.slides]
    assert titles == ["Kickoff", "Goals", "Budget"]
    assert "Ship v1" in prs.slides[1].placeholders[1].text_frame.text
    assert prs.slides[1].notes_slide.notes_text_frame.text == "say hi"

    from_blocks = Presentation(io.BytesIO(build_file("pptx", DOC_SPEC).data))
    assert [s.shapes.title.text for s in from_blocks.slides][:2] == ["Q3 report", "Summary"]


@pytest.mark.parametrize("fmt, spec, message", [
    ("exe", {}, "format must be one of"),
    ("docx", [], "spec must be an object"),
    ("docx", {}, "needs 'blocks'"),
    ("docx", {"blocks": [{"type": "script"}]}, "type must be"),
    ("xlsx", {"rows": "a,b"}, "rows must be a list of lists"),
    ("xlsx", {"rows": [[1]] * 100_001}, "more than 100000 rows"),
    ("pptx", {"slides": "x"}, "'slides' must be a list"),
])
def test_bad_specs_explain_what_to_fix(fmt, spec, message):
    with pytest.raises(FileSpecError, match=message):
        build_file(fmt, spec)


def test_safe_filename():
    assert safe_filename("../../etc/passwd", "pdf") == "etcpasswd.pdf"
    assert safe_filename("Q3 report.xlsx", "xlsx") == "Q3 report.xlsx"
    assert safe_filename("", "docx") == "document.docx"


def test_create_owned_file_stores_it_as_the_owners_upload(tmp_path, monkeypatch):
    import src.tool_utils as tool_utils
    from src.file_store import create_owned_file
    from src.upload_handler import UploadHandler

    handler = UploadHandler(base_dir=str(tmp_path), upload_dir=str(tmp_path / "uploads"))
    monkeypatch.setattr(tool_utils, "_upload_handler", handler)
    out = create_owned_file("xlsx", {"rows": [["a"], [1]]}, "sales", "alice")
    assert out["name"] == "sales.xlsx" and out["url"] == f"/api/upload/{out['id']}"
    info = next(i for i in handler._load_upload_index().values() if i["id"] == out["id"])
    assert info["owner"] == "alice"


def test_create_owned_file_without_storage_fails_clearly(monkeypatch):
    import src.tool_utils as tool_utils
    from src.file_store import FileStoreUnavailable, create_owned_file

    monkeypatch.setattr(tool_utils, "_upload_handler", None)
    with pytest.raises(FileStoreUnavailable):
        create_owned_file("csv", {"rows": [["a"]]}, "x", "alice")


# ---------------------------------------------------------------------------
# create_file agent tool
# ---------------------------------------------------------------------------

def _tool_with_store(tmp_path, monkeypatch):
    import src.tool_utils as tool_utils
    from src.agent_tools.file_tools import CreateFileTool
    from src.upload_handler import UploadHandler

    handler = UploadHandler(base_dir=str(tmp_path), upload_dir=str(tmp_path / "uploads"))
    monkeypatch.setattr(tool_utils, "_upload_handler", handler)
    return CreateFileTool(), handler


def test_create_file_tool_with_nested_spec(tmp_path, monkeypatch):
    import asyncio, json

    tool, handler = _tool_with_store(tmp_path, monkeypatch)
    out = asyncio.run(tool.execute(json.dumps({
        "format": "xlsx", "filename": "Q3 sales",
        "spec": {"sheets": [{"name": "Sales", "rows": [["Region", "Revenue"], ["North", 120]]}]},
    }), {"owner": "alice"}))
    assert out["exit_code"] == 0
    assert out["file"]["name"] == "Q3_sales.xlsx"  # the upload store turns spaces into _
    assert out["file"]["url"] in out["response"]
    info = next(i for i in handler._load_upload_index().values() if i["id"] == out["file"]["id"])
    assert info["owner"] == "alice"


def test_create_file_tool_accepts_top_level_fields(tmp_path, monkeypatch):
    import asyncio

    tool, _ = _tool_with_store(tmp_path, monkeypatch)
    out = asyncio.run(tool.execute({"format": "pdf", "title": "Memo",
                                    "blocks": [{"type": "paragraph", "text": "Hi"}]}, {"owner": "bob"}))
    assert out["exit_code"] == 0 and out["file"]["name"].endswith(".pdf")


def test_create_file_tool_reports_bad_specs(tmp_path, monkeypatch):
    import asyncio

    tool, _ = _tool_with_store(tmp_path, monkeypatch)
    out = asyncio.run(tool.execute({"format": "docx", "spec": {"blocks": "nope"}}, {"owner": "bob"}))
    assert out["exit_code"] == 1 and "Could not create the file" in out["error"]
    out = asyncio.run(tool.execute("{not json", {"owner": "bob"}))
    assert out["exit_code"] == 1
