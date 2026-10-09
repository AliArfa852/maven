"""Charts in created files (plan v2 §7): native in xlsx/pptx/pdf, an image in docx."""

import io

import pytest

from src import file_charts
from src.file_builder import FileSpecError, build_file

CHART = {"type": "chart", "kind": "bar", "title": "Revenue by region",
         "categories": ["North", "South", "East"],
         "series": [{"name": "2025", "values": [120, 80, 95]},
                    {"name": "2026", "values": [140, "85", None]}]}


@pytest.mark.parametrize("kind", file_charts.KINDS)
def test_png_renders_every_kind(kind):
    from PIL import Image

    png = file_charts.render_png(file_charts.parse({**CHART, "kind": kind}))
    img = Image.open(io.BytesIO(png))
    assert img.format == "PNG" and img.size == (1400, 800)
    # Something other than white was drawn.
    assert len(img.convert("RGB").getcolors(1_000_000)) > 5


def test_png_handles_negative_values_and_many_categories():
    spec = {"kind": "bar", "categories": [f"d{i}" for i in range(200)],
            "series": [{"name": "p&l", "values": [(-1) ** i * i for i in range(200)]}]}
    assert file_charts.render_png(file_charts.parse(spec))[:4] == b"\x89PNG"
    flat = {"kind": "line", "categories": ["a", "b"], "series": [{"values": [0, 0]}]}
    assert file_charts.render_png(file_charts.parse(flat))[:4] == b"\x89PNG"


@pytest.mark.parametrize("bad, msg", [
    ({"kind": "radar"}, "kind must be"),
    ({"categories": []}, "categories"),
    ({"series": [{"values": [1, 2]}]}, "2 values for 3 categories"),
    ({"series": [{"values": [1, "lots", 3]}]}, "not a number"),
    ({"kind": "pie", "series": [{"values": [1, -2, 3]}]}, "negative"),
    ({"series": [{"values": [1, 2, 3]}] * 13}, "more than 12 series"),
])
def test_bad_chart_specs_explain_themselves(bad, msg):
    with pytest.raises(FileSpecError, match=msg):
        build_file("docx", {"blocks": [{**CHART, **bad}]})


def test_docx_embeds_the_chart_as_a_picture():
    from docx import Document

    doc = Document(io.BytesIO(build_file("docx", {"title": "R", "blocks": [CHART]}).data))
    assert len(doc.inline_shapes) == 1


def test_pdf_draws_the_chart():
    from pypdf import PdfReader

    data = build_file("pdf", {"title": "R", "blocks": [CHART, {**CHART, "kind": "pie"},
                                                       {**CHART, "kind": "line"},
                                                       {**CHART, "kind": "barh"}]}).data
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(data)).pages)
    assert "North" in text and "2026" in text


def test_pptx_chart_slide_is_a_native_chart():
    from pptx import Presentation

    data = build_file("pptx", {"slides": [{"title": "Revenue", "chart": CHART},
                                          {"title": "Mix", "chart": {**CHART, "kind": "pie"}}]}).data
    prs = Presentation(io.BytesIO(data))
    charts = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
    assert len(charts) == 2
    assert list(charts[0].plots[0].categories) == ["North", "South", "East"]
    assert [s.name for s in charts[0].plots[0].series] == ["2025", "2026"]
    assert list(charts[0].plots[0].series[1].values) == [140, 85, None]


def test_blocks_to_slides_turn_a_chart_block_into_a_slide():
    from pptx import Presentation

    data = build_file("pptx", {"blocks": [{"type": "heading", "text": "Sales"}, CHART]}).data
    prs = Presentation(io.BytesIO(data))
    assert sum(sh.has_chart for s in prs.slides for sh in s.shapes) == 1


def test_xlsx_chart_draws_from_sheet_columns():
    from openpyxl import load_workbook

    rows = [["Region", "2025", "2026"], ["North", 120, 140], ["South", 80, 85]]
    data = build_file("xlsx", {"sheets": [{"name": "Sales", "rows": rows, "charts": [
        {"kind": "bar", "title": "Revenue"},
        {"kind": "pie", "y": [2], "anchor": "H30"},
    ]}]}).data
    ws = load_workbook(io.BytesIO(data))["Sales"]
    assert len(ws._charts) == 2
    assert ws._charts[0].series[0].val.numRef.f == "'Sales'!$B$1:$B$3" or \
        "Sales" in ws._charts[0].series[0].val.numRef.f


@pytest.mark.parametrize("chart, msg", [
    ({"y": [7]}, "column numbers"),
    ({"x": "A"}, "column numbers"),
    ({"anchor": "=A1"}, "anchor"),
])
def test_xlsx_chart_errors(chart, msg):
    rows = [["Region", "Revenue"], ["North", 1]]
    with pytest.raises(FileSpecError, match=msg):
        build_file("xlsx", {"sheets": [{"rows": rows, "charts": [chart]}]})


def test_xlsx_chart_needs_data():
    with pytest.raises(FileSpecError, match="header row and data"):
        build_file("xlsx", {"sheets": [{"rows": [["a", "b"]], "charts": [{}]}]})
