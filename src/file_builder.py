"""Build office files from a simple structured spec (plan v2 §5.1, R12).

One spec shape for every format, so an agent or the UI describes content once:

    {"title": "Q3 report",
     "blocks": [
        {"type": "heading", "text": "Summary", "level": 1},
        {"type": "paragraph", "text": "Revenue grew 12%."},
        {"type": "bullets", "items": ["North +20%", "South +4%"]},
        {"type": "table", "rows": [["Region", "Revenue"], ["North", 120]]},
        {"type": "chart", "kind": "bar", "categories": ["North", "South"],
         "series": [{"name": "Revenue", "values": [120, 80]}]}
     ]}

Spreadsheets take ``{"sheets": [{"name": "Data", "rows": [[...], ...]}]}``
(``rows`` alone means one sheet); the first row is the header, and a sheet's
``charts`` list draws from its columns. Slides take ``{"slides": [{"title": "...",
"bullets": [...]} | {"title": "...", "table": rows} | {"title": "...", "chart": {...}}]}``,
and ``blocks`` work too (one slide per heading). Chart specs: src/file_charts.py.

build_file() returns the bytes; it never touches disk or the network.
Libraries (all MIT/BSD): openpyxl, python-docx, python-pptx, reportlab.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any

from src import file_charts
from src.file_charts import ChartSpecError

FORMATS = {
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "pdf": "application/pdf",
    "csv": "text/csv",
}

# Generous limits: a runaway agent should not build a gigabyte file.
MAX_ROWS = 100_000
MAX_COLS = 200
MAX_BLOCKS = 5_000
MAX_SLIDES = 300
MAX_TEXT = 100_000


class FileSpecError(ValueError):
    """The spec is malformed; the message says what to fix."""


@dataclass
class BuiltFile:
    data: bytes
    mime: str
    extension: str


def safe_filename(name: str, extension: str) -> str:
    stem = re.sub(r"[^\w\- .]+", "", str(name or "")).strip(" .")[:80] or "document"
    stem = re.sub(r"\.(xlsx|docx|pptx|pdf|csv)$", "", stem, flags=re.IGNORECASE)
    return f"{stem}.{extension}"


# ── spec helpers ───────────────────────────────────────────────────────────

def _text(value: Any) -> str:
    s = "" if value is None else str(value)
    if len(s) > MAX_TEXT:
        raise FileSpecError(f"text longer than {MAX_TEXT} characters")
    return s


def _rows(rows: Any, where: str) -> list[list[Any]]:
    if not isinstance(rows, list) or not all(isinstance(r, list) for r in rows):
        raise FileSpecError(f"{where}: rows must be a list of lists")
    if len(rows) > MAX_ROWS:
        raise FileSpecError(f"{where}: more than {MAX_ROWS} rows")
    if any(len(r) > MAX_COLS for r in rows):
        raise FileSpecError(f"{where}: more than {MAX_COLS} columns")
    return rows


def _cell(value: Any) -> Any:
    """Numbers and booleans stay typed; everything else becomes text."""
    if value is None or isinstance(value, (int, float, bool)):
        return value
    return _text(value)


def _blocks(spec: dict) -> list[dict]:
    blocks = spec.get("blocks")
    if blocks is None:
        if "text" in spec:
            blocks = [{"type": "paragraph", "text": spec["text"]}]
        elif "rows" in spec:
            blocks = [{"type": "table", "rows": spec["rows"]}]
        else:
            raise FileSpecError("document needs 'blocks' (or 'text' / 'rows')")
    if not isinstance(blocks, list) or len(blocks) > MAX_BLOCKS:
        raise FileSpecError(f"'blocks' must be a list of at most {MAX_BLOCKS} items")
    out = []
    for i, b in enumerate(blocks):
        if not isinstance(b, dict) or b.get("type") not in {"heading", "paragraph", "bullets", "table", "chart"}:
            raise FileSpecError(f"block {i}: type must be heading, paragraph, bullets, table or chart")
        out.append(b)
    return out


def _sheets(spec: dict) -> list[tuple[str, list[list[Any]], list]]:
    sheets = spec.get("sheets")
    if sheets is None:
        if "rows" not in spec:
            raise FileSpecError("spreadsheet needs 'sheets' or 'rows'")
        sheets = [{"name": spec.get("title") or "Sheet1", "rows": spec["rows"]}]
    if not isinstance(sheets, list) or not sheets:
        raise FileSpecError("'sheets' must be a non-empty list")
    out, seen = [], set()
    for i, sh in enumerate(sheets):
        if not isinstance(sh, dict):
            raise FileSpecError(f"sheet {i}: must be an object with name and rows")
        # Excel: max 31 chars, no []:*?/\ , unique names.
        name = re.sub(r"[\[\]:*?/\\]", "", _text(sh.get("name") or f"Sheet{i + 1}"))[:31] or f"Sheet{i + 1}"
        base, n = name, 2
        while name.lower() in seen:
            name = f"{base[:28]}_{n}"
            n += 1
        seen.add(name.lower())
        charts = sh.get("charts") or []
        if not isinstance(charts, list) or len(charts) > file_charts.MAX_CHARTS_PER_SHEET:
            raise FileSpecError(f"sheet '{name}': 'charts' must be a list of at most "
                                f"{file_charts.MAX_CHARTS_PER_SHEET}")
        out.append((name, _rows(sh.get("rows", []), f"sheet '{name}'"), charts))
    return out


# ── formats ────────────────────────────────────────────────────────────────

def _xlsx(spec: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows, charts in _sheets(spec):
        ws = wb.create_sheet(title=name)
        for r in rows:
            ws.append([_formula_safe(_cell(v)) for v in r])
        if rows:
            for c in ws[1]:
                c.font = Font(bold=True)
            ws.freeze_panes = "A2"
            for idx in range(1, max(len(r) for r in rows) + 1):
                width = max((len(str(r[idx - 1])) for r in rows if len(r) >= idx and r[idx - 1] is not None),
                            default=8)
                ws.column_dimensions[get_column_letter(idx)].width = min(max(width + 2, 8), 60)
        for i, c in enumerate(charts):
            where = f"sheet '{name}' chart {i}"
            if len(rows) < 2:
                raise FileSpecError(f"{where}: the sheet needs a header row and data")
            kind, x, y = file_charts.sheet_columns(c, rows[0], where)
            anchor = c.get("anchor") or f"{get_column_letter(len(rows[0]) + 2)}{2 + 20 * i}"
            if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", str(anchor)):
                raise FileSpecError(f"{where}: anchor must be a cell like H2")
            file_charts.add_xlsx_chart(ws, kind, x, y, len(rows), _text(c.get("title")), anchor)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _formula_safe(value: Any) -> Any:
    """Text starting with = + - @ would run as a formula when opened (CSV/formula
    injection); prefix it so it stays text. Numbers are untouched."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _csv(spec: dict) -> bytes:
    rows = _rows(spec.get("rows") if "rows" in spec else _sheets(spec)[0][1], "csv")
    buf = io.StringIO()
    writer = csv.writer(buf)
    for r in rows:
        writer.writerow(["" if v is None else _formula_safe(_cell(v)) for v in r])
    # BOM so Excel opens UTF-8 (names with accents) correctly.
    return ("﻿" + buf.getvalue()).encode("utf-8")


def _docx(spec: dict) -> bytes:
    from docx import Document

    doc = Document()
    if spec.get("title"):
        doc.add_heading(_text(spec["title"]), level=0)
    for b in _blocks(spec):
        kind = b["type"]
        if kind == "heading":
            doc.add_heading(_text(b.get("text")), level=min(max(int(b.get("level", 1)), 1), 4))
        elif kind == "paragraph":
            doc.add_paragraph(_text(b.get("text")))
        elif kind == "bullets":
            for item in b.get("items") or []:
                doc.add_paragraph(_text(item), style="List Bullet")
        elif kind == "table":
            rows = _rows(b.get("rows", []), "table")
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            table.style = "Table Grid"
            for i, r in enumerate(rows):
                for j in range(cols):
                    table.cell(i, j).text = "" if j >= len(r) or r[j] is None else _text(r[j])
            for c in table.rows[0].cells:
                for run in c.paragraphs[0].runs:
                    run.bold = True
        elif kind == "chart":
            from docx.shared import Inches

            png = file_charts.render_png(file_charts.parse(b, "chart"))
            doc.add_picture(io.BytesIO(png), width=Inches(6.3))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf(spec: dict) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    flow = []
    if spec.get("title"):
        flow += [Paragraph(escape(_text(spec["title"])), styles["Title"]), Spacer(1, 4 * mm)]
    for b in _blocks(spec):
        kind = b["type"]
        if kind == "heading":
            level = min(max(int(b.get("level", 1)), 1), 3)
            flow.append(Paragraph(escape(_text(b.get("text"))), styles[f"Heading{level}"]))
        elif kind == "paragraph":
            for para in _text(b.get("text")).split("\n\n"):
                flow.append(Paragraph(escape(para).replace("\n", "<br/>"), styles["BodyText"]))
        elif kind == "bullets":
            items = [ListItem(Paragraph(escape(_text(i)), styles["BodyText"])) for i in b.get("items") or []]
            if items:
                flow.append(ListFlowable(items, bulletType="bullet"))
        elif kind == "table":
            rows = _rows(b.get("rows", []), "table")
            if not rows:
                continue
            cols = max(len(r) for r in rows)
            cell = styles["BodyText"]
            data = [[Paragraph(escape("" if j >= len(r) or r[j] is None else _text(r[j])), cell)
                     for j in range(cols)] for r in rows]
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf2")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            flow.append(t)
        elif kind == "chart":
            flow.append(file_charts.pdf_drawing(file_charts.parse(b, "chart"), 170 * mm, 95 * mm))
        flow.append(Spacer(1, 3 * mm))
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, title=_text(spec.get("title") or ""),
                      leftMargin=18 * mm, rightMargin=18 * mm,
                      topMargin=18 * mm, bottomMargin=18 * mm).build(flow or [Spacer(1, 1)])
    return buf.getvalue()


def _slides_from(spec: dict) -> list[dict]:
    slides = spec.get("slides")
    if slides is not None:
        if not isinstance(slides, list):
            raise FileSpecError("'slides' must be a list")
        return slides
    # Blocks -> one slide per heading, bullets/paragraphs as points.
    out: list[dict] = []
    for b in _blocks(spec):
        if b["type"] == "heading" or not out:
            out.append({"title": _text(b.get("text")) if b["type"] == "heading" else "", "bullets": []})
            if b["type"] == "heading":
                continue
        if b["type"] == "paragraph":
            out[-1]["bullets"].append(_text(b.get("text")))
        elif b["type"] == "bullets":
            out[-1]["bullets"].extend(_text(i) for i in b.get("items") or [])
        elif b["type"] == "table":
            out.append({"title": out[-1]["title"], "table": b.get("rows", [])})
        elif b["type"] == "chart":
            out.append({"title": b.get("title") or out[-1]["title"], "chart": b})
    return out


def _pptx(spec: dict) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)  # 16:9
    slides = _slides_from(spec)
    if len(slides) > MAX_SLIDES:
        raise FileSpecError(f"more than {MAX_SLIDES} slides")
    if spec.get("title"):
        s = prs.slides.add_slide(prs.slide_layouts[0])
        s.shapes.title.text = _text(spec["title"])
        if spec.get("subtitle") and len(s.placeholders) > 1:
            s.placeholders[1].text = _text(spec["subtitle"])
    for i, sd in enumerate(slides):
        if not isinstance(sd, dict):
            raise FileSpecError(f"slide {i}: must be an object")
        if "chart" in sd:
            chart = file_charts.parse(sd["chart"], f"slide {i} chart")
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = _text(sd.get("title") or chart.title)
            file_charts.add_pptx_chart(s, chart, Inches(0.8), Inches(1.5),
                                       prs.slide_width - Inches(1.6), prs.slide_height - Inches(2.0))
        elif "table" in sd:
            rows = _rows(sd["table"], f"slide {i} table")
            s = prs.slides.add_slide(prs.slide_layouts[5])
            s.shapes.title.text = _text(sd.get("title"))
            if rows:
                cols = max(len(r) for r in rows)
                shape = s.shapes.add_table(len(rows), cols, Inches(0.6), Inches(1.5),
                                           prs.slide_width - Inches(1.2), Inches(0.4) * len(rows))
                for r_i, r in enumerate(rows):
                    for c_i in range(cols):
                        cell = shape.table.cell(r_i, c_i)
                        cell.text = "" if c_i >= len(r) or r[c_i] is None else _text(r[c_i])
                        for p in cell.text_frame.paragraphs:
                            for run in p.runs:
                                run.font.size = Pt(14)
        else:
            s = prs.slides.add_slide(prs.slide_layouts[1])
            s.shapes.title.text = _text(sd.get("title"))
            body = s.placeholders[1].text_frame
            body.clear()
            for j, point in enumerate(sd.get("bullets") or []):
                p = body.paragraphs[0] if j == 0 else body.add_paragraph()
                p.text = _text(point)
        if sd.get("notes"):
            s.notes_slide.notes_text_frame.text = _text(sd["notes"])
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


_BUILDERS = {"xlsx": _xlsx, "csv": _csv, "docx": _docx, "pdf": _pdf, "pptx": _pptx}


def build_file(fmt: str, spec: dict) -> BuiltFile:
    """Build a file of ``fmt`` from ``spec``. Raises FileSpecError on bad input."""
    fmt = str(fmt or "").strip().lower().lstrip(".")
    if fmt not in _BUILDERS:
        raise FileSpecError(f"format must be one of: {', '.join(sorted(_BUILDERS))}")
    if not isinstance(spec, dict):
        raise FileSpecError("spec must be an object")
    try:
        data = _BUILDERS[fmt](spec)
    except FileSpecError:
        raise
    except ChartSpecError as e:
        raise FileSpecError(str(e)) from e
    except ImportError as e:
        raise FileSpecError(f"{fmt} support is not installed ({e.name}); "
                            "run pip install -r requirements.txt") from e
    except (TypeError, ValueError, KeyError) as e:
        raise FileSpecError(f"could not build {fmt}: {e}") from e
    return BuiltFile(data=data, mime=FORMATS[fmt], extension=fmt)
