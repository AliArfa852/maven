"""Exact answers from a spreadsheet or CSV (plan v2 §5.2, first step).

The model is bad at adding up a column it has only half seen. This module
loads the user's own Excel/CSV file and computes the numbers: describe,
filter, group-by with aggregates (month/quarter/year buckets for dates) and
top-N. No code execution: a fixed set of operations over parsed cells, so
it is safe for every employee (the admin-only Python tool stays separate).

Numbers in business files are messy; "1,234.50", "1.234,50", "$1,200",
"(500)" (an accounting negative) and "12%" are read as numbers.
"""
from __future__ import annotations

import csv
import io
import math
import re
import statistics
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

MAX_ROWS = 200_000
MAX_OUT_ROWS = 200
AGGREGATES = ("sum", "mean", "count", "min", "max", "median")
OPS = ("=", "!=", ">", ">=", "<", "<=", "contains", "not contains")
BUCKETS = ("day", "month", "quarter", "year")


class AnalysisError(ValueError):
    """The request or the file cannot be analysed; the message says why."""


@dataclass
class Table:
    columns: list[str]
    rows: list[list[Any]]
    sheet: str | None = None
    sheets: list[str] = field(default_factory=list)
    truncated: bool = False


# ── reading ────────────────────────────────────────────────────────────────

def load_table(path: str, name: str = "", sheet: str | None = None, header_row: int = 1) -> Table:
    lower = (name or path).lower()
    if lower.endswith((".xlsx", ".xlsm")):
        return _load_xlsx(path, sheet, header_row)
    if lower.endswith((".csv", ".tsv", ".txt")):
        return _load_csv(path, header_row)
    raise AnalysisError("Only Excel (.xlsx) and CSV/TSV files can be analysed")


def _load_xlsx(path: str, sheet: str | None, header_row: int) -> Table:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = wb.sheetnames
        if sheet:
            match = next((n for n in names if n.lower() == str(sheet).strip().lower()), None)
            if match is None:
                raise AnalysisError(f"No sheet named {sheet!r}. Sheets: {', '.join(names)}")
        else:
            match = names[0]
        rows_iter = wb[match].iter_rows(values_only=True)
        return _table(rows_iter, header_row, sheet=match, sheets=names)
    finally:
        wb.close()


def _load_csv(path: str, header_row: int) -> Table:
    with open(path, "rb") as f:
        raw = f.read(64 * 1024 * 1024)
    text = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    # Sniff the delimiter from the header row on: a title line above it fools the sniffer.
    lines = text.splitlines()
    sample = "\n".join(lines[max(0, int(header_row or 1) - 1):][:200])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _table(csv.reader(io.StringIO(text), dialect), header_row)


def _table(rows_iter, header_row: int, sheet=None, sheets=None) -> Table:
    header_row = max(1, int(header_row or 1))
    header = None
    rows: list[list[Any]] = []
    truncated = False
    for i, row in enumerate(rows_iter, 1):
        if i < header_row:
            continue
        row = list(row)
        if header is None:
            header = row
            continue
        if not any(v not in (None, "") for v in row):
            continue
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
        rows.append(row)
    if header is None:
        raise AnalysisError("The sheet is empty")
    columns, seen = [], {}
    for j, h in enumerate(header):
        label = str(h).strip() if h not in (None, "") else f"Column {j + 1}"
        if label.lower() in seen:
            seen[label.lower()] += 1
            label = f"{label} ({seen[label.lower()]})"
        else:
            seen[label.lower()] = 1
        columns.append(label)
    width = len(columns)
    rows = [(r + [None] * width)[:width] for r in rows]
    # Trim trailing empty columns Excel often reports.
    while columns and columns[-1].startswith("Column ") and all(r[len(columns) - 1] in (None, "") for r in rows):
        columns.pop()
        rows = [r[:len(columns)] for r in rows]
    return Table(columns, rows, sheet, list(sheets or []), truncated)


# ── values ─────────────────────────────────────────────────────────────────

_NUM_STRIP = re.compile(r"[\s$€£¥₹]|(?i:usd|eur|gbp|pkr|rs\.?)")


def to_number(v: Any) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return None if isinstance(v, float) and not math.isfinite(v) else float(v)
    s = str(v).strip()
    if not s:
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    pct = s.endswith("%")
    s = _NUM_STRIP.sub("", s.rstrip("%"))
    if s.startswith("-"):
        neg, s = not neg, s[1:]
    if not s or not re.fullmatch(r"[\d.,]+", s):
        return None
    if "," in s and "." in s:
        # The later separator is the decimal one: 1,234.5 or 1.234,5
        s = s.replace(",", "") if s.rfind(".") > s.rfind(",") else s.replace(".", "").replace(",", ".")
    elif "," in s:
        parts = s.split(",")
        s = s.replace(",", "") if all(len(p) == 3 for p in parts[1:]) else s.replace(",", ".")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    try:
        n = float(s)
    except ValueError:
        return None
    n = -n if neg else n
    return n / 100 if pct else n


def to_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%d.%m.%Y",
                    "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%b %Y", "%B %Y", "%Y-%m"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def _bucket(v: Any, by: str) -> Any:
    d = to_date(v)
    if d is None:
        return None if v in (None, "") else str(v)
    if by == "day":
        return d.isoformat()
    if by == "month":
        return f"{d.year}-{d.month:02d}"
    if by == "quarter":
        return f"{d.year}-Q{(d.month - 1) // 3 + 1}"
    return str(d.year)


def _kind(values: list[Any]) -> str:
    present = [v for v in values if v not in (None, "")]
    if not present:
        return "empty"
    sample = present[:500]
    if sum(to_number(v) is not None for v in sample) >= 0.9 * len(sample):
        return "number"
    if sum(to_date(v) is not None for v in sample) >= 0.9 * len(sample):
        return "date"
    return "text"


def _fmt(v: Any) -> Any:
    if isinstance(v, float):
        if v.is_integer() and abs(v) < 1e15:
            return int(v)
        return round(v, 4)
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return v


# ── operations ─────────────────────────────────────────────────────────────

def _col(table: Table, name: Any) -> int:
    if isinstance(name, int) and 0 <= name < len(table.columns):
        return name
    want = str(name or "").strip().lower()
    for i, c in enumerate(table.columns):
        if c.lower() == want:
            return i
    raise AnalysisError(f"No column named {name!r}. Columns: {', '.join(table.columns)}")


def _matches(cell: Any, op: str, value: Any) -> bool:
    if op in ("contains", "not contains"):
        hit = str(value).lower() in ("" if cell is None else str(cell)).lower()
        return hit if op == "contains" else not hit
    a, b = to_number(cell), to_number(value)
    if a is None or b is None:
        da, db = to_date(cell), to_date(value)
        if da and db:
            a, b = da, db
        else:
            a = "" if cell is None else str(cell).strip().lower()
            b = "" if value is None else str(value).strip().lower()
            if op not in ("=", "!="):
                return False
    return {"=": a == b, "!=": a != b, ">": a > b, ">=": a >= b, "<": a < b, "<=": a <= b}[op]


def apply_filters(table: Table, where: Any) -> list[list[Any]]:
    if not where:
        return table.rows
    if isinstance(where, dict):
        where = [where]
    if not isinstance(where, list):
        raise AnalysisError("'where' must be a list of {column, op, value}")
    conds = []
    for w in where:
        if not isinstance(w, dict):
            raise AnalysisError("each 'where' item must be {column, op, value}")
        op = str(w.get("op") or "=").strip().lower()
        if op == "==":
            op = "="
        if op not in OPS:
            raise AnalysisError(f"op must be one of: {', '.join(OPS)}")
        conds.append((_col(table, w.get("column")), op, w.get("value")))
    return [r for r in table.rows if all(_matches(r[i], op, v) for i, op, v in conds)]


def describe(table: Table, rows: list[list[Any]]) -> dict:
    cols = []
    for i, name in enumerate(table.columns):
        values = [r[i] for r in rows]
        kind = _kind(values)
        info: dict[str, Any] = {"column": name, "type": kind,
                                "filled": sum(v not in (None, "") for v in values),
                                "empty": sum(v in (None, "") for v in values)}
        if kind == "number":
            nums = [n for n in (to_number(v) for v in values) if n is not None]
            info.update(sum=_fmt(sum(nums)), mean=_fmt(statistics.fmean(nums)),
                        median=_fmt(statistics.median(nums)), min=_fmt(min(nums)), max=_fmt(max(nums)))
        elif kind == "date":
            ds = [d for d in (to_date(v) for v in values) if d]
            info.update(first=ds and min(ds).isoformat(), last=ds and max(ds).isoformat())
        else:
            counts: dict[str, int] = {}
            for v in values:
                if v not in (None, ""):
                    counts[str(v)] = counts.get(str(v), 0) + 1
            info["distinct"] = len(counts)
            info["most_common"] = sorted(counts.items(), key=lambda kv: (-kv[1], kv[0]))[:5]
        cols.append(info)
    return {"rows": len(rows), "columns": cols}


def _aggregate(fn: str, values: list[Any]) -> Any:
    if fn == "count":
        return sum(v not in (None, "") for v in values)
    nums = [n for n in (to_number(v) for v in values) if n is not None]
    if not nums:
        return None
    return _fmt({"sum": math.fsum(nums), "mean": statistics.fmean(nums), "min": min(nums),
                 "max": max(nums), "median": statistics.median(nums)}[fn])


def group(table: Table, rows: list[list[Any]], group_by: Any, aggregates: Any,
          sort: Any = None, descending: bool = True, limit: int = 50) -> dict:
    if isinstance(group_by, (str, int)):
        group_by = [group_by]
    if not isinstance(group_by, list) or not group_by:
        raise AnalysisError("'group_by' needs one or more columns (add ':month', ':quarter' or ':year' to bucket a date column)")
    keys = []
    for g in group_by:
        name, _, by = str(g).partition(":")
        by = by.strip().lower()
        if by and by not in BUCKETS:
            raise AnalysisError(f"date bucket must be one of: {', '.join(BUCKETS)}")
        keys.append((_col(table, name.strip()), by, f"{table.columns[_col(table, name.strip())]}{' (' + by + ')' if by else ''}"))
    if isinstance(aggregates, dict):
        aggregates = [aggregates]
    aggregates = aggregates or [{"fn": "count"}]
    aggs = []
    for a in aggregates:
        if not isinstance(a, dict):
            raise AnalysisError("each aggregate must be {column, fn}")
        fn = str(a.get("fn") or a.get("function") or "sum").strip().lower()
        if fn in ("avg", "average"):
            fn = "mean"
        if fn not in AGGREGATES:
            raise AnalysisError(f"fn must be one of: {', '.join(AGGREGATES)}")
        if fn == "count" and a.get("column") in (None, ""):
            aggs.append((None, fn, "count"))
        else:
            ci = _col(table, a.get("column"))
            aggs.append((ci, fn, f"{fn} of {table.columns[ci]}"))

    buckets: dict[tuple, list[list[Any]]] = {}
    for r in rows:
        key = tuple(_bucket(r[ci], by) if by else (_fmt(r[ci]) if r[ci] not in ("",) else None)
                    for ci, by, _ in keys)
        buckets.setdefault(key, []).append(r)
    out_cols = [k[2] for k in keys] + [a[2] for a in aggs]
    out = []
    for key, members in buckets.items():
        vals = [len(members) if ci is None else _aggregate(fn, [m[ci] for m in members]) for ci, fn, _ in aggs]
        out.append(list(key) + vals)

    sort_idx = len(keys)  # first aggregate by default
    if sort not in (None, ""):
        names = [c.lower() for c in out_cols]
        if str(sort).lower() not in names:
            raise AnalysisError(f"sort must be one of: {', '.join(out_cols)}")
        sort_idx = names.index(str(sort).lower())
    if any(k[1] for k in keys) and sort in (None, ""):
        out.sort(key=lambda r: tuple((v is None, "" if v is None else str(v)) for v in r[:len(keys)]))  # time order, blanks last
    else:
        def order(r):
            v = r[sort_idx]
            return (0, v) if isinstance(v, (int, float)) else (1, str(v))

        present = sorted((r for r in out if r[sort_idx] is not None), key=order, reverse=bool(descending))
        out = present + [r for r in out if r[sort_idx] is None]  # missing values last
    total = len(out)
    return {"columns": out_cols, "rows": out[: _limit(limit)], "groups": total}


def top(table: Table, rows: list[list[Any]], sort_by: Any, descending: bool = True,
        limit: int = 10, columns: Any = None) -> dict:
    ci = _col(table, sort_by)
    picked = [_col(table, c) for c in columns] if columns else list(range(len(table.columns)))
    keyed = [(to_number(r[ci]), r) for r in rows]
    keyed = [kv for kv in keyed if kv[0] is not None]
    keyed.sort(key=lambda kv: kv[0], reverse=bool(descending))
    return {"columns": [table.columns[i] for i in picked],
            "rows": [[_fmt(r[i]) for i in picked] for _, r in keyed[: _limit(limit)]],
            "matching_rows": len(keyed)}


def _limit(limit: Any) -> int:
    try:
        return max(1, min(int(limit or 50), MAX_OUT_ROWS))
    except (TypeError, ValueError):
        return 50


def run(table: Table, request: dict) -> dict:
    """Apply one analysis request: {operation, where?, ...operation fields}."""
    op = str(request.get("operation") or "describe").strip().lower()
    rows = apply_filters(table, request.get("where"))
    base = {"sheet": table.sheet, "sheets": table.sheets, "rows_in_file": len(table.rows),
            "rows_after_filter": len(rows), "file_truncated": table.truncated}
    if op == "describe":
        return {**base, "operation": op, **describe(table, rows)}
    if op in ("group", "group_by", "pivot", "summarize"):
        return {**base, "operation": "group", **group(
            table, rows, request.get("group_by"), request.get("aggregates"),
            request.get("sort"), request.get("descending", True), request.get("limit", 50))}
    if op in ("top", "sort"):
        return {**base, "operation": "top", **top(
            table, rows, request.get("sort_by"), request.get("descending", True),
            request.get("limit", 10), request.get("columns"))}
    if op in ("rows", "filter"):
        cols = request.get("columns")
        picked = [_col(table, c) for c in cols] if cols else list(range(len(table.columns)))
        return {**base, "operation": "rows", "columns": [table.columns[i] for i in picked],
                "rows": [[_fmt(r[i]) for i in picked] for r in rows[: _limit(request.get("limit", 50))]]}
    raise AnalysisError("operation must be describe, group, top or rows")


def to_markdown(result: dict) -> str:
    head = (f"Sheet {result['sheet']!r}: " if result.get("sheet") else "") + \
        f"{result['rows_after_filter']:,} of {result['rows_in_file']:,} rows"
    if result.get("file_truncated"):
        head += f" (file has more than {MAX_ROWS:,} rows; only the first {MAX_ROWS:,} were read)"
    if result["operation"] == "describe":
        lines = [head, "", "| column | type | filled | summary |", "| --- | --- | --- | --- |"]
        for c in result["columns"]:
            if c["type"] == "number":
                s = f"sum {c['sum']}, mean {c['mean']}, median {c['median']}, min {c['min']}, max {c['max']}"
            elif c["type"] == "date":
                s = f"{c['first']} to {c['last']}"
            elif c["type"] == "text":
                s = f"{c['distinct']} distinct; most common: " + ", ".join(f"{k} ({n})" for k, n in c["most_common"])
            else:
                s = ""
            lines.append(f"| {c['column']} | {c['type']} | {c['filled']} | {s} |")
        return "\n".join(lines)
    cols, rows = result["columns"], result["rows"]

    def cell(v):
        return "" if v is None else str(v).replace("|", "\\|").replace("\n", " ")

    lines = [head + (f"; {result['groups']:,} groups" if "groups" in result else ""), "",
             "| " + " | ".join(cell(c) for c in cols) + " |", "|" + " --- |" * len(cols)]
    lines += ["| " + " | ".join(cell(v) for v in r) + " |" for r in rows]
    shown = len(rows)
    total = result.get("groups", result.get("matching_rows", result.get("rows_after_filter")))
    if total and total > shown:
        lines.append(f"\n(showing {shown} of {total:,})")
    return "\n".join(lines)
